"""Oylik hisobotni Excel (.xlsx) va PDF qilib tayyorlash.

collect()   -- bazadan ma'lumot yig'adi (bir joyda, ikkala format uchun bir xil raqamlar)
build_xlsx() / build_pdf() -- tayyor fayl baytlarini qaytaradi
PDF uchun DejaVu shrifti ishlatiladi (apps/finance/fonts): kirill va o'zbek harflari to'g'ri chiqadi.
"""
import calendar
import io
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from django.db.models import Q, Sum
from django.utils import timezone

from apps.animals.models import AnimalGroup
from apps.inventory.models import InventoryItem

from .models import Transaction

MONTHS = [
    "Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun",
    "Iyul", "Avgust", "Sentabr", "Oktabr", "Noyabr", "Dekabr",
]
FONT_DIR = Path(__file__).resolve().parent / 'fonts'
GREEN = '25402F'

CATEGORY_LABELS = dict(Transaction.CATEGORY_CHOICES)


def month_label(year, month):
    return f"{MONTHS[month - 1]} {year}"


def money(value):
    """12 345 678 ko'rinishida (probel bilan)."""
    return f"{Decimal(value or 0):,.0f}".replace(',', ' ')


def qty(value):
    return format(Decimal(value).normalize(), 'f')


def person(user):
    if not user:
        return ''
    return user.get_full_name() or user.username


def _totals(farm, start, end):
    row = Transaction.objects.filter(farm=farm, date__range=(start, end)).aggregate(
        income=Sum('amount', filter=Q(transaction_type='kirim')),
        expense=Sum('amount', filter=Q(transaction_type='chiqim')),
    )
    income, expense = row['income'] or Decimal(0), row['expense'] or Decimal(0)
    return {'income': income, 'expense': expense, 'profit': income - expense}


def collect(farm, year, month, generated_by=None):
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    prev_end = start - timedelta(days=1)
    prev_start = prev_end.replace(day=1)

    txs = list(
        Transaction.objects.filter(farm=farm, date__range=(start, end))
        .select_related('recorded_by').order_by('date', 'id')
    )

    categories = {'kirim': [], 'chiqim': []}
    for kind in categories:
        rows = (
            Transaction.objects.filter(farm=farm, date__range=(start, end), transaction_type=kind)
            .values('category').annotate(total=Sum('amount')).order_by('-total')
        )
        categories[kind] = [(CATEGORY_LABELS.get(r['category'], r['category']), r['total']) for r in rows]

    items = list(InventoryItem.objects.filter(farm=farm).order_by('item_name'))
    animals = list(
        AnimalGroup.objects.filter(farm=farm, status='faol')
        .values('animal_type__name').annotate(total=Sum('count')).order_by('-total')
    )

    return {
        'farm': farm,
        'year': year,
        'month': month,
        'label': month_label(year, month),
        'start': start,
        'end': end,
        'generated_at': timezone.localtime(),
        'generated_by': person(generated_by),
        'totals': _totals(farm, start, end),
        'prev_label': month_label(prev_start.year, prev_start.month),
        'prev_totals': _totals(farm, prev_start, prev_end),
        'categories': categories,
        'milk': _milk_summary(farm, start, end),
        'transactions': [
            {
                'date': t.date,
                'type': t.get_transaction_type_display(),
                'is_income': t.transaction_type == 'kirim',
                'category': t.get_category_display(),
                'amount': t.amount,
                'description': t.description,
                'who': person(t.recorded_by),
            }
            for t in txs
        ],
        'inventory': [
            {
                'name': i.item_name,
                'quantity': i.quantity,
                'unit': i.unit,
                'low': i.quantity <= i.low_stock_threshold,
            }
            for i in items
        ],
        'animals': [(a['animal_type__name'], a['total']) for a in animals],
    }


def _milk_summary(farm, start, end):
    """Oy bo'yicha sut: sog'ilgan, sotilgan, narx, tannarx, yemga ketgan xarajat va zararli sigirlar soni."""
    from apps.animals import profit as milk_profit
    from apps.animals.models import MilkSale

    today = timezone.localdate()
    if start > today:
        return None
    rep = milk_profit.build_report(farm, start, min(end, today))
    t = rep['totals']
    if Decimal(t['liters']) == 0 and not MilkSale.objects.filter(farm=farm, date__range=(start, end)).exists():
        return None
    sold = MilkSale.objects.filter(farm=farm, date__range=(start, end)).aggregate(l=Sum('liters'), a=Sum('amount'))
    return {
        'liters': Decimal(t['liters']),
        'sold_liters': sold['l'] or Decimal('0'),
        'sold_amount': sold['a'] or Decimal('0'),
        'price': Decimal(t['price_per_liter']) if t['price_per_liter'] else None,
        'cost_per_liter': Decimal(t['cost_per_liter']) if t['cost_per_liter'] else None,
        'feed_cost': Decimal(t['feed_cost']),
        'feed_per_liter': Decimal(t['feed_cost_per_liter']) if t['feed_cost_per_liter'] else None,
        'profit': Decimal(t['profit']) if t['profit'] is not None else None,
        'cows': t['cows'],
        'losing_cows': t['losing_cows'],
    }


def _milk_rows(m):
    """Sut bo'limi qatorlari (PDF va Excel uchun bir xil)."""
    def som(v):
        return f"{money(v)} so'm" if v is not None else '—'
    return [
        ("Sog'ilgan sut", f"{qty(m['liters'])} l"),
        ('Sotilgan sut', f"{qty(m['sold_liters'])} l ({som(m['sold_amount'])})"),
        ('1 litr sotuv narxi', som(m['price'])),
        ('1 litr sut tannarxi (hamma xarajat)', som(m['cost_per_liter'])),
        ('Shundan yemga (1 litr)', som(m['feed_per_liter'])),
        ('Yem xarajati (oy)', som(m['feed_cost'])),
        ('Sutdan foyda / zarar', som(m['profit'])),
        ('Zarar keltirayotgan sigirlar', f"{m['losing_cows']} / {m['cows']}"),
    ]


def result_word(profit):
    return "Foyda" if profit >= 0 else "Zarar"


# ---------------------------------------------------------------- Excel

def build_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill('solid', fgColor=GREEN)
    head_font = Font(bold=True, color='FFFFFF')
    bold = Font(bold=True)
    money_fmt = '#,##0'

    def header(ws, row, titles):
        for col, title in enumerate(titles, 1):
            c = ws.cell(row=row, column=col, value=title)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal='center', vertical='center')

    def widths(ws, values):
        for i, w in enumerate(values, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # 1-varaq: xulosa
    ws = wb.active
    ws.title = 'Xulosa'
    t, p = data['totals'], data['prev_totals']
    ws['A1'] = f"{data['farm'].name}: {data['label']} hisoboti"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Manzil: {data['farm'].location}"
    ws['A3'] = f"Tayyorlandi: {data['generated_at']:%d.%m.%Y %H:%M}" + (
        f" ({data['generated_by']})" if data['generated_by'] else '')

    header(ws, 5, ['Ko\'rsatkich', data['label'], data['prev_label']])
    rows = [
        ('Jami kirim', t['income'], p['income']),
        ('Jami chiqim', t['expense'], p['expense']),
        (f"Sof natija ({result_word(t['profit']).lower()})", t['profit'], p['profit']),
    ]
    for i, (name, cur, prev) in enumerate(rows, 6):
        ws.cell(row=i, column=1, value=name).font = bold if i == 8 else Font()
        for col, v in ((2, cur), (3, prev)):
            c = ws.cell(row=i, column=col, value=float(v))
            c.number_format = money_fmt
            if i == 8:
                c.font = Font(bold=True, color='1B7A3B' if v >= 0 else 'B3261E')

    r = 10
    for kind, title in (('kirim', 'Kirim kategoriyalari'), ('chiqim', 'Chiqim kategoriyalari')):
        header(ws, r, [title, 'Summa (so\'m)'])
        r += 1
        if not data['categories'][kind]:
            ws.cell(row=r, column=1, value="Yozuv yo'q")
            r += 1
        for name, total in data['categories'][kind]:
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=float(total)).number_format = money_fmt
            r += 1
        r += 1

    if data.get('milk'):
        header(ws, r, ['Sut', 'Qiymat'])
        r += 1
        for name, value in _milk_rows(data['milk']):
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=value)
            r += 1
        r += 1

    if data['animals']:
        header(ws, r, ['Hayvonlar (faol)', 'Soni'])
        r += 1
        for name, total in data['animals']:
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=total)
            r += 1
    widths(ws, [34, 20, 20])

    # 2-varaq: kirim-chiqim
    ws2 = wb.create_sheet('Kirim-chiqim')
    header(ws2, 1, ['Sana', 'Turi', 'Kategoriya', 'Summa (so\'m)', 'Izoh', 'Kiritgan'])
    for i, tx in enumerate(data['transactions'], 2):
        ws2.cell(row=i, column=1, value=tx['date']).number_format = 'DD.MM.YYYY'
        ws2.cell(row=i, column=2, value=tx['type'])
        ws2.cell(row=i, column=3, value=tx['category'])
        c = ws2.cell(row=i, column=4, value=float(tx['amount']))
        c.number_format = money_fmt
        c.font = Font(color='1B7A3B' if tx['is_income'] else 'B3261E')
        ws2.cell(row=i, column=5, value=tx['description'])
        ws2.cell(row=i, column=6, value=tx['who'])
    ws2.freeze_panes = 'A2'
    if data['transactions']:
        ws2.auto_filter.ref = f"A1:F{len(data['transactions']) + 1}"
    widths(ws2, [12, 10, 20, 18, 44, 22])

    # 3-varaq: ombor
    ws3 = wb.create_sheet('Ombor')
    header(ws3, 1, ['Mahsulot', 'Qoldiq', 'Birlik', 'Holat'])
    for i, it in enumerate(data['inventory'], 2):
        ws3.cell(row=i, column=1, value=it['name'])
        ws3.cell(row=i, column=2, value=float(it['quantity']))
        ws3.cell(row=i, column=3, value=it['unit'])
        c = ws3.cell(row=i, column=4, value='Kam qolgan' if it['low'] else 'Yetarli')
        if it['low']:
            c.font = Font(bold=True, color='B3261E')
    ws3.freeze_panes = 'A2'
    widths(ws3, [32, 14, 10, 16])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def build_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if 'DejaVu' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('DejaVu', str(FONT_DIR / 'DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('DejaVu-Bold', str(FONT_DIR / 'DejaVuSans-Bold.ttf')))
        pdfmetrics.registerFontFamily('DejaVu', normal='DejaVu', bold='DejaVu-Bold')

    green = colors.HexColor('#' + GREEN)
    red = colors.HexColor('#B3261E')
    ok = colors.HexColor('#1B7A3B')
    grid = colors.HexColor('#D8D8D0')

    h1 = ParagraphStyle('h1', fontName='DejaVu-Bold', fontSize=16, leading=20, textColor=green)
    h2 = ParagraphStyle('h2', fontName='DejaVu-Bold', fontSize=11, leading=14, spaceBefore=10, spaceAfter=4,
                        textColor=green)
    small = ParagraphStyle('small', fontName='DejaVu', fontSize=8, leading=10, textColor=colors.HexColor('#555555'))
    cell = ParagraphStyle('cell', fontName='DejaVu', fontSize=8, leading=10)

    def table(rows, col_widths, right_cols=(), bold_last=False):
        t = Table(rows, colWidths=col_widths, repeatRows=1)
        style = [
            ('FONT', (0, 0), (-1, -1), 'DejaVu', 8),
            ('FONT', (0, 0), (-1, 0), 'DejaVu-Bold', 8),
            ('BACKGROUND', (0, 0), (-1, 0), green),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.4, grid),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F6F6F1')]),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]
        for c in right_cols:
            style.append(('ALIGN', (c, 0), (c, -1), 'RIGHT'))
        if bold_last:
            style.append(('FONT', (0, -1), (-1, -1), 'DejaVu-Bold', 8))
        t.setStyle(TableStyle(style))
        return t

    farm, t, p = data['farm'], data['totals'], data['prev_totals']
    story = [
        Paragraph(f"{farm.name}: {data['label']} hisoboti", h1),
        Spacer(1, 2),
        Paragraph(
            f"Manzil: {farm.location} · Tayyorlandi: {data['generated_at']:%d.%m.%Y %H:%M}"
            + (f" · {data['generated_by']}" if data['generated_by'] else ''),
            small,
        ),
        Spacer(1, 8),
    ]

    summary = [
        ["Ko'rsatkich", data['label'], data['prev_label']],
        ['Jami kirim', f"{money(t['income'])} so'm", f"{money(p['income'])} so'm"],
        ['Jami chiqim', f"{money(t['expense'])} so'm", f"{money(p['expense'])} so'm"],
        [f"Sof natija ({result_word(t['profit']).lower()})", f"{money(t['profit'])} so'm",
         f"{money(p['profit'])} so'm"],
    ]
    st = table(summary, [70 * mm, 55 * mm, 55 * mm], right_cols=(1, 2), bold_last=True)
    st.setStyle(TableStyle([('TEXTCOLOR', (1, 3), (1, 3), ok if t['profit'] >= 0 else red)]))
    story.append(st)

    if data.get('milk'):
        story.append(Paragraph('Sut', h2))
        story.append(table([['Ko\'rsatkich', 'Qiymat']] + [list(r) for r in _milk_rows(data['milk'])],
                           [110 * mm, 70 * mm], right_cols=(1,)))

    for kind, title in (('kirim', 'Kirim kategoriyalari'), ('chiqim', 'Chiqim kategoriyalari')):
        story.append(Paragraph(title, h2))
        rows = [['Kategoriya', "Summa (so'm)"]] + [[n, money(v)] for n, v in data['categories'][kind]]
        if len(rows) == 1:
            story.append(Paragraph("Bu oyda yozuv yo'q", small))
        else:
            story.append(table(rows, [110 * mm, 70 * mm], right_cols=(1,)))

    story.append(Paragraph("Kirim-chiqim ro'yxati", h2))
    if data['transactions']:
        rows = [['Sana', 'Turi', 'Kategoriya', "Summa (so'm)", 'Izoh']]
        for tx in data['transactions']:
            rows.append([
                f"{tx['date']:%d.%m.%Y}", tx['type'], tx['category'], money(tx['amount']),
                Paragraph((tx['description'] or '').replace('&', '&amp;').replace('<', '&lt;'), cell),
            ])
        story.append(table(rows, [22 * mm, 16 * mm, 32 * mm, 30 * mm, 80 * mm], right_cols=(3,)))
    else:
        story.append(Paragraph("Bu oyda yozuv yo'q", small))

    story.append(Paragraph('Ombor holati', h2))
    if data['inventory']:
        rows = [['Mahsulot', 'Qoldiq', 'Holat']] + [
            [i['name'], f"{qty(i['quantity'])} {i['unit']}", 'Kam qolgan' if i['low'] else 'Yetarli']
            for i in data['inventory']
        ]
        it = table(rows, [90 * mm, 50 * mm, 40 * mm], right_cols=(1,))
        low_rows = [n for n, i in enumerate(data['inventory'], 1) if i['low']]
        it.setStyle(TableStyle([('TEXTCOLOR', (2, n), (2, n), red) for n in low_rows]))
        story.append(it)
    else:
        story.append(Paragraph("Omborda mahsulot yo'q", small))

    if data['animals']:
        story.append(Paragraph('Hayvonlar (faol)', h2))
        story.append(table([['Turi', 'Soni']] + [[n, str(c)] for n, c in data['animals']],
                           [110 * mm, 70 * mm], right_cols=(1,)))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('DejaVu', 7)
        canvas.setFillColor(colors.HexColor('#777777'))
        canvas.drawString(15 * mm, 8 * mm, 'FarmLine')
        canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"{doc.page}-bet")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
        title=f"{farm.name} {data['label']}", author='FarmLine',
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
