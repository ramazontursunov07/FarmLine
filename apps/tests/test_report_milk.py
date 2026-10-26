import io
from datetime import timedelta
from decimal import Decimal

from openpyxl import load_workbook

from apps.animals.models import MilkRecord, MilkSale
from apps.finance.models import Transaction
from apps.tests.test_animals import AnimalCardBase

EXPORT = '/api/finance/export/'


class MonthlyReportMilkTests(AnimalCardBase):
    def setUp(self):
        super().setUp()
        self.month = self.today.strftime('%Y-%m')
        cow = self.make_animal('1', birth_date=self.today - timedelta(days=1500))
        MilkRecord.objects.create(animal=cow, date=self.today, session='ertalab', liters=Decimal('40'),
                                  recorded_by=self.owner)
        MilkSale.objects.create(farm=self.farm, date=self.today, liters=Decimal('30'),
                                price_per_liter=Decimal('5000'), amount=Decimal('150000'), recorded_by=self.owner)
        Transaction.objects.create(farm=self.farm, transaction_type='chiqim', category='ozuqa',
                                   amount=Decimal('200000'), date=self.today, recorded_by=self.owner)

    def labels(self):
        r = self.client.get(EXPORT, {'farm': self.farm.pk, 'month': self.month, 'file': 'xlsx'})
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content))['Xulosa']
        return {row[0].value: row[1].value for row in ws.iter_rows(min_row=1) if row[0].value}

    def test_xlsx_has_milk_section(self):
        rows = self.labels()
        self.assertEqual(rows["Sog'ilgan sut"], '40 l')
        self.assertIn('30 l', rows['Sotilgan sut'])
        self.assertIn('5 000', rows['1 litr sotuv narxi'])
        self.assertIn('5 000', rows['Shundan yemga (1 litr)'])  # 200000 / 40
        self.assertIn('200 000', rows['Yem xarajati (oy)'])

    def test_pdf_with_milk(self):
        r = self.client.get(EXPORT, {'farm': self.farm.pk, 'month': self.month, 'file': 'pdf'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_month_without_milk_has_no_section(self):
        r = self.client.get(EXPORT, {'farm': self.farm.pk, 'month': '2020-01', 'file': 'xlsx'})
        ws = load_workbook(io.BytesIO(r.content))['Xulosa']
        self.assertNotIn("Sog'ilgan sut", {row[0].value for row in ws.iter_rows()})
