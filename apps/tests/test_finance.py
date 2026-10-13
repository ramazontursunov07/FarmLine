from rest_framework.test import APITestCase

from apps.farms.models import Farm
from apps.users.models import User


class FinanceFilterTests(APITestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='fermer1', password='Kuchli-Parol-2026', role='fermer')
        Farm.objects.create(owner=self.owner, name='Baraka', location='Qarshi', farm_type='aralash')
        self.client.force_authenticate(self.owner)

    def test_bad_dates_return_400_not_500(self):
        for params in ('date_from=2026-13-45', 'date_to=kecha', 'date_from=abc'):
            r = self.client.get(f'/api/finance/?{params}')
            self.assertEqual(r.status_code, 400, params)

    def test_bad_farm_returns_400(self):
        self.assertEqual(self.client.get('/api/finance/?farm=abc').status_code, 400)

    def test_valid_filters_ok(self):
        r = self.client.get('/api/finance/?date_from=2026-01-01&date_to=2026-12-31')
        self.assertEqual(r.status_code, 200)


import io
from datetime import date

from openpyxl import load_workbook

from apps.farms.models import Worker
from apps.finance.models import Transaction
from apps.inventory.models import InventoryItem

EXPORT = '/api/finance/export/'


def make_user(username, role='fermer'):
    return User.objects.create_user(username=username, password='Kuchli-Parol-2026', role=role)


class ExportTests(APITestCase):
    def setUp(self):
        self.owner = make_user('egasi')
        self.farm = Farm.objects.create(owner=self.owner, name="Baraka ferma", location='Qarshi', farm_type='aralash')
        for kind, cat, amount, d in (
            ('kirim', 'sotuv', 1000000, date(2026, 10, 2)),
            ('chiqim', 'ozuqa', 1500000, date(2026, 10, 5)),
            ('chiqim', 'dori', 200000, date(2026, 9, 28)),
        ):
            Transaction.objects.create(farm=self.farm, transaction_type=kind, category=cat, amount=amount,
                                       date=d, recorded_by=self.owner, description="Izoh <b>&</b>")
        InventoryItem.objects.create(farm=self.farm, item_name='Omuxta yem', item_type='ozuqa', quantity=3,
                                     unit='kg', low_stock_threshold=10)
        self.client.force_authenticate(self.owner)

    def get(self, **params):
        params.setdefault('farm', self.farm.pk)
        return self.client.get(EXPORT, params)

    def test_xlsx_numbers(self):
        r = self.get(month='2026-10', file='xlsx')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn(f'farmline-{self.farm.pk}-2026-10.xlsx', r['Content-Disposition'])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Xulosa', 'Kirim-chiqim', 'Ombor'])
        summary = wb['Xulosa']
        self.assertEqual(summary['B6'].value, 1000000)   # kirim
        self.assertEqual(summary['B7'].value, 1500000)   # chiqim
        self.assertEqual(summary['B8'].value, -500000)   # zarar
        self.assertEqual(summary['C7'].value, 200000)    # o'tgan oy chiqimi
        self.assertEqual(wb['Kirim-chiqim'].max_row, 3)  # sarlavha + 2 ta (sentabr kirmaydi)
        self.assertEqual(wb['Ombor']['D2'].value, 'Kam qolgan')

    def test_pdf_ok(self):
        r = self.get(month='2026-10', file='pdf')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_cyrillic_and_empty_month(self):
        self.farm.name = 'Ферма «Барака»'
        self.farm.save()
        r = self.get(month='2025-01', file='pdf')   # yanvar: o'tgan oy = dekabr (oldingi yil)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.get(month='2025-01', file='xlsx').status_code, 200)

    def test_defaults_to_current_month_xlsx(self):
        r = self.get()
        self.assertEqual(r.status_code, 200)
        self.assertIn('.xlsx', r['Content-Disposition'])

    def test_validation_errors(self):
        self.assertEqual(self.get(month='2026-13').status_code, 400)
        self.assertEqual(self.get(month='oktabr').status_code, 400)
        self.assertEqual(self.get(file='doc').status_code, 400)
        self.assertEqual(self.client.get(EXPORT).status_code, 400)
        self.assertEqual(self.client.get(EXPORT, {'farm': 'abc'}).status_code, 400)

    def test_other_farmer_cannot_export(self):
        self.client.force_authenticate(make_user('boshqa'))
        self.assertEqual(self.get(month='2026-10').status_code, 404)

    def test_worker_needs_finance_permission(self):
        worker = make_user('ishchi', 'ishchi')
        link = Worker.objects.create(user=worker, farm=self.farm, can_view_finance=False)
        self.client.force_authenticate(worker)
        self.assertEqual(self.get().status_code, 404)
        link.can_view_finance = True
        link.save()
        self.assertEqual(self.get().status_code, 200)

    def test_admin_can_export_any_farm(self):
        self.client.force_authenticate(make_user('admin1', 'admin'))
        self.assertEqual(self.get(file='pdf').status_code, 200)

    def test_requires_login(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.get().status_code, 401)
